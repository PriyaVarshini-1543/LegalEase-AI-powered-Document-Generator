from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from exporters.common import stream


def export_docx(title: str, content: str):
    document = Document()

    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for block in content.split("\n\n"):
        block = block.strip()
        if not block:
            continue

        lines = block.splitlines()
        if len(lines) == 1 and (
            lines[0].isupper() or lines[0].endswith(":")
        ):
            document.add_heading(lines[0].rstrip(":"), level=1)
        else:
            paragraph = document.add_paragraph()
            paragraph.add_run(block)

    output = BytesIO()
    document.save(output)
    return stream(output.getvalue())
